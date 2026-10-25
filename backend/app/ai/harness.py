"""
NexLab Unified Multi-Role AI Harness Runtime (`backend/app/ai/harness.py`).

Powered by the Claw-Code Autonomous Engine (`claw_runtime.py`):
1. Unified Autonomous Execution (no artificial barrier between Chat and Agent modes):
   - Auto-executes creation, setup, experiment+HTML+quiz authoring, enrollment,
     scheduling, grading, collection-building, playground-saving, and report tools
     directly in the database across up to `MAX_TURNS = 6` chained tool steps per turn!
   - Only pauses for confirmation on destructive operations (`CONFIRM_ACTION_NAMES`).
2. Built-in Slash Command Graph (`/setup`, `/experiment`, `/exam`, `/schedule`,
   `/autograde`, `/collection`, `/playground`, `/studyplan`, `/report`, `/onboard`).
3. 6-Layer Dynamic System Prompt Builder with 2-Stage Tool Pool Routing.
4. Rolling Memory Auto-Compaction (`maybe_compact_session_memory`).
5. Deterministic Code-Level Socratic Output Filter (`enforce_socratic_guardrail`) for
   `student` and `platform_user`.
"""
import csv
import io
import json
import logging
import re
from datetime import datetime

from flask import g
from sqlalchemy import or_

from app.extensions import db, socketio
from app.ai.gate import guarded_ai_call, AIDisabled, AIQuotaExceeded
from app.ai.agent import (
    AgentToolError,
    run_web_search,
    _extract_json,
    _truncate,
)
from app.ai.claw_runtime import (
    QueryEngineConfig,
    UsageSummary,
    PermissionDenial,
    SLASH_COMMANDS,
    match_slash_command,
)
from app.ai.role_tools import get_tools_for_role, CONFIRM_ACTION_NAMES
from app.models.user import User
from app.models.institution import Institution
from app.models.admin_copilot import AdminChatSession, AdminChatMessage, InstitutionDocument
from app.ai.rag_embeddings import get_embedding, cosine_sim
from app.models.ai_documents import DocumentChunk, HAS_PGVECTOR

log = logging.getLogger(__name__)

ENGINE_CONFIG = QueryEngineConfig(max_turns=6, max_budget_tokens=32000, compact_after_turns=16)


# ══════════════════════════════════════════════════════════════════════════════
# 1. Multi-Format Document Text & Chunk Extractor (.pdf, .docx, .xlsx, .csv, .md, .txt)
# ══════════════════════════════════════════════════════════════════════════════

SUPPORTED_RAG_EXTENSIONS = ('.pdf', '.docx', '.xlsx', '.csv', '.md', '.txt')


def extract_multi_format_chunks(file_path: str, filename: str, file_id: str) -> tuple[list[dict], int, str]:
    ext = ('.' + filename.rsplit('.', 1)[-1].lower()) if '.' in filename else '.txt'
    doc_type = ext.lstrip('.')
    pages_text: list[tuple[int, str]] = []

    if ext == '.pdf':
        import pypdf
        reader = pypdf.PdfReader(file_path)
        for p_idx, page in enumerate(reader.pages, start=1):
            t = (page.extract_text() or '').strip()
            if t:
                pages_text.append((p_idx, t))

    elif ext == '.docx':
        try:
            import docx
            doc = docx.Document(file_path)
            paras = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
            for tbl in doc.tables:
                for row in tbl.rows:
                    row_txt = ' | '.join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_txt:
                        paras.append(row_txt)
            full_txt = '\n'.join(paras)
            for idx in range(0, max(1, len(full_txt)), 2000):
                pages_text.append(((idx // 2000) + 1, full_txt[idx:idx + 2000]))
        except Exception as exc:
            log.warning("[harness.rag] python-docx fallback for %s: %s", filename, exc)
            import zipfile
            import xml.etree.ElementTree as ET
            with zipfile.ZipFile(file_path) as zf:
                xml_content = zf.read('word/document.xml')
            tree = ET.fromstring(xml_content)
            text_nodes = [elem.text for elem in tree.iter() if elem.text]
            full_txt = ' '.join(text_nodes)
            pages_text.append((1, full_txt))

    elif ext == '.xlsx':
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        for s_idx, sheet_name in enumerate(wb.sheetnames, start=1):
            ws = wb[sheet_name]
            lines = [f"Sheet: {sheet_name}"]
            for row in ws.iter_rows(values_only=True):
                vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if vals:
                    lines.append(' | '.join(vals))
            if len(lines) > 1:
                pages_text.append((s_idx, '\n'.join(lines)))

    elif ext == '.csv':
        with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
            reader = csv.reader(f)
            lines = [' | '.join(row) for row in reader if any(c.strip() for c in row)]
        full_txt = '\n'.join(lines)
        for idx in range(0, max(1, len(full_txt)), 2000):
            pages_text.append(((idx // 2000) + 1, full_txt[idx:idx + 2000]))

    else:
        with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
            full_txt = f.read().strip()
        for idx in range(0, max(1, len(full_txt)), 2000):
            pages_text.append(((idx // 2000) + 1, full_txt[idx:idx + 2000]))

    chunks = []
    step = 500
    for page_num, text in pages_text:
        for i in range(0, len(text), step):
            chunk_text = text[i:i + 650].strip()
            if len(chunk_text) > 25:
                chunks.append({
                    'chunk_id': f"{file_id}_p{page_num}_{len(chunks)}",
                    'page': page_num,
                    'text': chunk_text,
                    'tokens': len(chunk_text.split()),
                })

    return chunks, max(1, len(pages_text)), doc_type


def vector_search_scoped_documents(institution_id, user_id, role, query, doc_id=None, top_k=4):
    """Semantic RAG retrieval over DocumentChunk embeddings (pgvector or JSON-list cosine)."""
    if not query or not query.strip():
        return []

    q = InstitutionDocument.query
    if doc_id:
        q = q.filter(InstitutionDocument.id == doc_id)
    elif institution_id:
        q = q.filter(or_(
            InstitutionDocument.institution_id == institution_id,
            InstitutionDocument.uploaded_by == user_id,
        ))
    else:
        q = q.filter(InstitutionDocument.uploaded_by == user_id)
    docs = q.all()
    if not docs:
        return []

    doc_ids = [d.id for d in docs]
    chunk_rows = DocumentChunk.query.filter(DocumentChunk.document_id.in_(doc_ids)).all()
    if not chunk_rows:
        return []

    query_emb = get_embedding(query)
    doc_map = {d.id: d for d in docs}
    scored = []
    for ch in chunk_rows:
        if ch.embedding is None:
            continue
        if HAS_PGVECTOR:
            dist = ch.embedding.cosine_distance(query_emb) if hasattr(ch.embedding, 'cosine_distance') else None
            score = None
            # fall back to python cosine if pgvector object lacks method
            try:
                if dist is None:
                    score = cosine_sim(list(ch.embedding), query_emb)
                else:
                    score = 1.0 - float(dist)
            except Exception:
                score = None
        else:
            score = cosine_sim(ch.embedding, query_emb)
        if score is None:
            continue
        # keyword-overlap re-rank: penalise chunks sharing no tokens with query
        q_tokens = set(query.lower().split())
        ch_tokens = set((ch.text or '').lower().split())
        if not (q_tokens & ch_tokens):
            score *= 0.85
        scored.append({
            'doc_id': ch.document_id,
            'filename': doc_map[ch.document_id].filename if ch.document_id in doc_map else 'document',
            'page': ch.page or 1,
            'text': ch.text,
            'similarity_score': round(float(max(0.0, score)), 3),
        })

    scored.sort(key=lambda r: r['similarity_score'], reverse=True)
    return scored[:top_k]


def search_scoped_documents(institution_id: str, user_id: str, role: str,
                            query: str, doc_id: str = None, top_k: int = 4) -> list[dict]:
    if not query or not query.strip():
        return []

    q = InstitutionDocument.query
    if doc_id:
        q = q.filter(InstitutionDocument.id == doc_id)
    elif institution_id:
        q = q.filter(or_(
            InstitutionDocument.institution_id == institution_id,
            InstitutionDocument.uploaded_by == user_id,
        ))
    else:
        q = q.filter(InstitutionDocument.uploaded_by == user_id)

    docs = q.all()
    if not docs:
        return []

    # Pgvector / semantic path first, when available; TF-IDF is the guaranteed fallback.
    if HAS_PGVECTOR:
        try:
            has_chunks = DocumentChunk.query.filter(
                DocumentChunk.document_id.in_([d.id for d in docs])
            ).first() is not None
            if has_chunks:
                vec_results = vector_search_scoped_documents(
                    institution_id, user_id, role, query, doc_id=doc_id, top_k=top_k)
                if vec_results:
                    return vec_results
        except Exception as exc:
            log.warning("[harness.rag] vector search failed, falling back to TF-IDF: %s", exc)

    corpus_chunks = []
    for d in docs:
        for ch in (d.chunks or []):
            corpus_chunks.append({
                'doc_id': d.id,
                'filename': d.filename,
                'page': ch.get('page', 1),
                'text': ch.get('text', ''),
            })

    if not corpus_chunks:
        return []

    try:
        from sklearn.feature_extraction.text import TfidfVectorizer
        from sklearn.metrics.pairwise import cosine_similarity

        vectorizer = TfidfVectorizer(stop_words='english')
        chunk_texts = [c['text'] for c in corpus_chunks]
        tfidf_matrix = vectorizer.fit_transform(chunk_texts)
        query_vec = vectorizer.transform([query])
        scores = cosine_similarity(query_vec, tfidf_matrix).flatten()

        ranked_indices = scores.argsort()[::-1]
        results = []
        for idx in ranked_indices[:top_k]:
            if scores[idx] > 0.04:
                item = corpus_chunks[idx].copy()
                item['similarity_score'] = round(float(scores[idx]), 3)
                results.append(item)
        return results
    except Exception as exc:
        log.warning("[harness.rag] TF-IDF fallback: %s", exc)
        return corpus_chunks[:top_k]


# ══════════════════════════════════════════════════════════════════════════════
# 2. Deterministic Socratic Output Guardrail (Student & Platform-User Protection)
# ══════════════════════════════════════════════════════════════════════════════

_CODE_BLOCK_RE = re.compile(r"```(?:[a-zA-Z0-9_+-]*)\n(.*?)```", re.DOTALL)
_SOLUTION_SIGNATURE_RE = re.compile(
    r"^\s*(?:def\s+\w+\s*\(|class\s+\w+|#include\s*<|public\s+static\s+|int\s+main\s*\(|function\s+\w+\s*\()",
    re.MULTILINE,
)


def enforce_socratic_guardrail(text: str, role: str) -> tuple[str, bool]:
    if role not in ('student', 'platform_user') or not text:
        return text, False

    was_filtered = False

    def _replace_block(match: re.Match) -> str:
        nonlocal was_filtered
        code_body = (match.group(1) or '').strip()
        non_empty_lines = [
            ln for ln in code_body.splitlines()
            if ln.strip() and not ln.strip().startswith(('#', '//', '/*'))
        ]
        if len(non_empty_lines) > 6 or _SOLUTION_SIGNATURE_RE.search(code_body):
            was_filtered = True
            return (
                "> 🔒 **Socratic Tutor Guardrail:** Complete code solutions are hidden so you can build mastery! "
                "Try translating the step-by-step logic above into code yourself, then click **Run** or ask me to check your attempt."
            )
        return match.group(0)

    filtered = _CODE_BLOCK_RE.sub(_replace_block, text)
    return filtered, was_filtered


# ══════════════════════════════════════════════════════════════════════════════
# 3. 2-Stage Page & Intent-Routed Tool Catalog
# ══════════════════════════════════════════════════════════════════════════════

def select_routed_tool_catalog(role_tools: dict, role_key: str, mode: str,
                               prompt: str = '', page_context: dict = None,
                               max_catalog_tools: int = 18) -> str:
    """
    Scores and selects the top tools for the prompt, prioritizing high-power
    autonomous builders (`autonomous_*`, `build_*`, `create_*`) when action verbs appear.
    All non-destructive tools are available in both 'chat' and 'agent' modes!
    """
    candidates = []
    p_tokens = set(re.findall(r'[a-z0-9_]+', (prompt or '').lower()))
    route_tokens = set(re.findall(r'[a-z0-9_]+', ((page_context or {}).get('route') or '').lower()))
    query_tokens = p_tokens | route_tokens

    has_action_verb = any(
        v in p_tokens for v in (
            'create', 'build', 'setup', 'add', 'enroll', 'assign', 'schedule',
            'generate', 'make', 'auto', 'publish', 'html', 'simulation', 'quiz',
            'exam', 'collection', 'playlist', 'playground', 'plan', 'onboard'
        )
    )

    for t_name, (_fn, desc, _allowed_mode) in role_tools.items():
        desc_tokens = set(re.findall(r'[a-z0-9_]+', f"{t_name} {desc}".lower()))
        overlap = len(query_tokens & desc_tokens) * 2
        boost = 0
        if has_action_verb and any(p in t_name for p in ('autonomous_', 'build_', 'create_', 'enroll_', 'generate_')):
            boost = 6
        candidates.append((overlap + boost, t_name, desc))

    candidates.sort(key=lambda item: (-item[0], item[1]))
    selected = candidates[:max_catalog_tools]
    return "\n".join(f"- {desc}" for (_sc, _name, desc) in selected)


# ══════════════════════════════════════════════════════════════════════════════
# 4. 6-Layer Dynamic System Prompt Stack
# ══════════════════════════════════════════════════════════════════════════════

ROLE_PERSONAS = {
    'student': """LAYER 2 — ROLE PERSONA: STUDENT SOCRATIC TUTOR & AUTONOMOUS WORKSPACE COMPANION
- You are an active Socratic teaching assistant and workspace companion for a university student.
- AUTONOMOUS CAPABILITIES: You can directly create custom Practice Problem Collections (`build_personal_problem_collection`), save starter Playground coding workspaces (`create_saved_playground_workspace`), enroll the student in Study Plans (`enroll_and_build_study_plan`), generate downloadable `.md` study sheets (`generate_study_sheet`), and inspect their live code, errors, grades, attendance, and SPARK scores.
- THE SOCRATIC CONTRACT (MANDATORY HARD RULES):
  1. NEVER write complete working solution code (>5 lines) for an experiment, exam, or practice problem.
  2. NEVER reveal hidden test case inputs/outputs (`is_hidden=True`).
  3. Guide the student using progressive 3-level hints (Level 1 Conceptual -> Level 2 Structural -> Level 3 Diagnostic).""",

    'platform_user': """LAYER 2 — ROLE PERSONA: SELF-PACED SOCRATIC CODING COACH & WORKSPACE ARCHITECT
- You are an active coding mentor and workspace builder for an independent learner on NexLab Platform.
- AUTONOMOUS CAPABILITIES: You can directly create custom Problem Collections & bookmarks (`build_personal_problem_collection`), create saved Playground workspaces (`create_saved_playground_workspace`), enroll in Study Plans (`enroll_and_build_study_plan`), generate downloadable study roadmaps/cheat sheets, and analyze practice submissions.
- NEVER dump full copy-pasteable solution code (>5 lines); teach the pattern and let the learner write the code.""",

    'instructor': """LAYER 2 — ROLE PERSONA: INSTRUCTOR AUTONOMOUS AUTHORING & CLASSROOM HARNESS
- You are an autonomous teaching and lab authoring harness for a university instructor.
- AUTONOMOUS EXECUTION: When the instructor asks to create an experiment (with HTML simulation, theory, runnable code, quiz, viva, test cases, and SPARK checkpoints), create an exam, add an HTML/Markdown/Code block, schedule sessions, auto-grade submissions, control a live session, or grant extensions/interventions, CALL THE AUTONOMOUS TOOL DIRECTLY (`build_complete_experiment_suite`, `autonomous_add_experiment_content_or_html`, `build_complete_exam`, `autonomous_bulk_autograde`, `autonomous_live_session_conductor`, `autonomous_grant_extension_or_intervention`, `autonomous_schedule_sessions`). It executes immediately in the database without manual forms!""",

    'institution_admin': """LAYER 2 — ROLE PERSONA: INSTITUTION ADMIN AUTONOMOUS OPERATIONS HARNESS
- You are an autonomous campus operations and analytics harness for the Institution Administrator.
- AUTONOMOUS EXECUTION: When the admin asks to set up a campus/semester hierarchy, create departments/academic years/batches/sections/subjects, assign instructors, enroll students, schedule single or recurring weekly lab sessions, create users, update AI quotas, or broadcast announcements, CALL THE AUTONOMOUS TOOLS DIRECTLY (`autonomous_campus_setup`, `autonomous_create_academic_entity`, `autonomous_assign_and_enroll`, `autonomous_schedule_sessions`, `autonomous_manage_user_and_governance`, `autonomous_broadcast_announcement`). You can chain multiple tool calls in a single turn to complete complex multi-step setups end-to-end!""",

    'admin': """LAYER 2 — ROLE PERSONA: PLATFORM SUPER-ADMIN AUTONOMOUS GOVERNANCE HARNESS
- You are the cross-tenant platform operations harness for the NexLab Super Admin.
- AUTONOMOUS EXECUTION: Directly onboard new Institution tenants (`autonomous_onboard_institution`), create global Platform Courses or Coding Contests (`autonomous_create_platform_course_or_contest`), update tenant plans/quotas or PlatformConfig flags (`autonomous_update_tenant_or_platform_config`), and query cross-tenant telemetry.""",
}


def build_six_layer_prompt(role: str, mode: str, user: User = None,
                           inst: Institution = None, page_context: dict = None,
                           compact_memory: str = None, rag_chunks: list = None,
                           catalog_str: str = '', is_planner: bool = False,
                           role_context: str = '') -> str:
    role_key = (role or 'student').lower().strip()
    u_name = user.full_name if user else 'User'
    inst_name = inst.name if inst else ('NexLab Platform' if role_key == 'platform_user' else 'NexLab Platform')
    ctx_dict = page_context or {}

    layer_1 = (
        "LAYER 1 — BASE IDENTITY:\n"
        "You are NexLab AI Harness (Claw-Code Autonomous Architecture), the native intelligence and execution engine inside the NexLab Virtual Programming Lab Platform.\n"
        f"Current UTC Timestamp: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}."
    )

    layer_2 = ROLE_PERSONAS.get(role_key, ROLE_PERSONAS['student'])

    ctx_lines = [
        f"- Authenticated User: {u_name} (Role: {role_key})",
        f"- Institution / Tenant: {inst_name}",
    ]
    if ctx_dict:
        if ctx_dict.get('route'):
            ctx_lines.append(f"- Active UI Route: {ctx_dict.get('route')}")
        if ctx_dict.get('page_title'):
            ctx_lines.append(f"- Active Page: {ctx_dict.get('page_title')}")
        if ctx_dict.get('experiment_id'):
            ctx_lines.append(f"- Active Experiment ID: {ctx_dict.get('experiment_id')}")
        if ctx_dict.get('problem_slug'):
            ctx_lines.append(f"- Active Practice Problem: {ctx_dict.get('problem_slug')}")
    if compact_memory:
        ctx_lines.append(f"- Compacted Prior Conversation Memory:\n{compact_memory}")

    layer_3 = "LAYER 3 — LIVE CONTEXT & MEMORY:\n" + "\n".join(ctx_lines)
    if role_context:
        layer_3 += "\n\n" + role_context

    layer_4 = (
        "LAYER 4 — NEXLAB DOMAIN & METRIC DEFINITIONS:\n"
        "- SPARK Score (0-100): Weighted behavioral index = 0.30*Consistency + 0.25*Efficiency + 0.20*Independence + 0.15*Engagement + 0.10*Integrity.\n"
        "- Risk Bands: Critical (<40% or >=3 high proctor flags), At-Risk (40-55% or <75% attendance), Watch (55-65%), Normal (>=65%).\n"
        "- Hint Deductions (Institutional Labs): Level 1 = -3%, Level 2 = -7%, Level 3 = -12% (capped at -15%)."
    )

    layer_5 = ""
    if rag_chunks:
        chunk_lines = [
            f"[{c.get('filename')}, p.{c.get('page')}]: {c.get('text', '')[:350]}"
            for c in rag_chunks[:4]
        ]
        layer_5 = "LAYER 5 — UPLOADED DOCUMENT RAG CONTEXT:\n" + "\n\n".join(chunk_lines)

    if is_planner:
        layer_6 = (
            f"LAYER 6 — ROUTED AUTONOMOUS TOOLS FOR ROLE '{role_key}':\n{catalog_str}\n\n"
            "CLAW-CODE QUERY ENGINE OUTPUT FORMAT:\n"
            "Reply with ONE strict JSON object and nothing else:\n"
            '  {"tool": "<tool_name>", "args": {...}}  — to execute a tool now\n'
            '  {"tool": null, "reason": "<all requested steps completed>"}\n'
            "Rules:\n"
            "- You can chain up to 6 tool calls sequentially! If the user asks for multiple tasks, execute them step-by-step until every requested item is created or queried.\n"
            "- For full campus/semester setups, prefer `autonomous_campus_setup`.\n"
            "- For creating an experiment or topic (including HTML simulation, theory, quiz, viva, test cases, checkpoints), prefer `build_complete_experiment_suite`.\n"
            "- For multi-page algorithm suites (such as Sorting or Searching under Data Structures), call `build_complete_experiment_suite` with `algorithm_pages=[\"Bubble Sort\", \"Selection Sort\", \"Insertion Sort\", \"Merge Sort\", \"Quick Sort\"]`. NEVER output massive raw HTML blocks in tool args; the system automatically builds the full standalone interactive visualizers and theory natively!\n"
            "- Never invent tool names or fabricate database IDs."
        )
    else:
        layer_6 = (
            "LAYER 6 — RESPONSE SYNTHESIS GUIDELINES:\n"
            "- Clearly confirm every database record, experiment, HTML simulation, quiz, section, subject, enrollment, session, playlist, or playground workspace that was created or queried.\n"
            "- Ground every figure, ID, URL, grade, or attendance percentage strictly in the tool observations provided.\n"
            "- Use clean Markdown formatting (bullet points, bold labels, compact tables when comparing items).\n"
            "- If a downloadable document (`generated_document`) was created, mention its title so the user can click the download card below."
        )

    parts = [layer_1, layer_2, layer_3, layer_4]
    if layer_5:
        parts.append(layer_5)
    parts.append(layer_6)
    return "\n\n".join(parts)


# ══════════════════════════════════════════════════════════════════════════════
# 5. Rolling Memory Compaction
# ══════════════════════════════════════════════════════════════════════════════

def maybe_compact_session_memory(session: AdminChatSession, institution_id: str, user_id: str, role: str):
    try:
        total_msgs = session.messages.count()
        if total_msgs < ENGINE_CONFIG.compact_after_turns or (total_msgs % 10 != 0 and session.compact_memory):
            return

        older_msgs = session.messages.order_by(AdminChatMessage.created_at.asc()).limit(max(1, total_msgs - 6)).all()
        if not older_msgs:
            return

        transcript = "\n".join(
            f"{m.role.upper()}: {(m.content or '')[:300]}"
            for m in older_msgs[-14:]
        )
        prompt_msgs = [
            {
                'role': 'system',
                'content': (
                    'Summarize the key entities created, user goals, resolved questions, and context from this '
                    'conversation transcript into a concise bulleted memory block under 180 words.'
                ),
            },
            {
                'role': 'user',
                'content': f"Prior Memory:\n{session.compact_memory or 'None'}\n\nRecent Turns:\n{transcript}",
            },
        ]
        summary_text, _, is_fb, _ = guarded_ai_call(
            'chat',
            institution_id=institution_id,
            user_id=user_id,
            role=role,
            messages=prompt_msgs,
            pool='fast',
            max_tokens=240,
            temperature=0.2,
        )
        if summary_text and not is_fb:
            session.compact_memory = summary_text.strip()[:1500]
            db.session.flush()
    except Exception as exc:
        log.warning("[harness.compaction] non-fatal memory compaction error: %s", exc)


def _emit_ai_event(user_id: str, event_name: str, payload: dict):
    if not user_id:
        return
    try:
        socketio.emit(event_name, payload, room=f"user_{user_id}")
    except Exception:
        pass


# ══════════════════════════════════════════════════════════════════════════════
# 6. Deterministic Autonomous Intent Fallback (When LLM is rate-limited or skips tool)
# ══════════════════════════════════════════════════════════════════════════════

def _extract_quoted_or_tail(prompt: str, default_val: str) -> str:
    q_match = re.search(r'["\']([^"\']{2,100})["\']', prompt)
    if q_match:
        return q_match.group(1).strip()
    m = re.search(r'(?:titled|named|called|on|for|about)\s+([A-Za-z0-9 &+_-]{3,70})', prompt, re.IGNORECASE)
    if m:
        return m.group(1).strip().rstrip('.,')
    return default_val


def _run_deterministic_autonomous_fallback(prompt: str, role_key: str, role_tools: dict) -> list[dict]:
    """
    Guarantees that if the LLM planner is rate-limited or fails to emit a JSON tool call,
    explicit user requests to create/setup/build/query still execute deterministically!
    """
    p_lower = prompt.lower()
    executed = []

    def _call(t_name: str, **kwargs):
        if t_name in role_tools:
            try:
                res = role_tools[t_name][0](**kwargs)
                executed.append({'tool': t_name, 'args': kwargs, 'result': res})
            except Exception as exc:
                executed.append({'tool': t_name, 'args': kwargs, 'error': str(exc)})

    if role_key == 'instructor':
        if any(w in p_lower for w in ('sort', 'sorting', 'dorting')) and any(w in p_lower for w in ('creat', 'crea', 'build', 'topic', 'page', 'experiment', 'simulat', 'datastruc')):
            _call('build_complete_experiment_suite', title='Sorting Algorithms', subject_name='Data Structures', algorithm_pages=['Bubble Sort', 'Selection Sort', 'Insertion Sort', 'Merge Sort', 'Quick Sort'])
        elif any(w in p_lower for w in ('create experiment', 'build experiment', 'new experiment', '/experiment', 'experiment on', 'experiment with', 'crea experiment', 'crea teh', 'create topic')):
            title = _extract_quoted_or_tail(prompt, 'Binary Search Tree & Traversal Lab')
            _call('build_complete_experiment_suite', title=title)
        elif any(w in p_lower for w in ('add html', 'simulation', 'visualizer', 'add module', 'add page')):
            _call('autonomous_add_experiment_content_or_html', module_type='html')
        elif any(w in p_lower for w in ('create exam', 'new exam', 'build exam', '/exam')):
            title = _extract_quoted_or_tail(prompt, 'Mid-Semester Coding & MCQ Exam')
            _call('build_complete_exam', title=title)
        elif any(w in p_lower for w in ('autograde', 'auto-grade', 'auto grade', '/autograde')):
            _call('autonomous_bulk_autograde', publish_immediately=True)
        elif any(w in p_lower for w in ('section', 'class', 'assigned')):
            _call('get_my_sections')
        elif any(w in p_lower for w in ('submission', 'plagiarism', 'ungraded')):
            _call('query_submissions_and_plagiarism')
        elif any(w in p_lower for w in ('error', 'mistake', 'bug', 'fail')):
            _call('get_common_errors')
        else:
            _call('get_section_analytics')

    elif role_key == 'institution_admin':
        if any(w in p_lower for w in ('/setup', 'full setup', 'campus setup', 'semester setup')) or (
            'department' in p_lower and ('section' in p_lower or 'subject' in p_lower or 'batch' in p_lower)
        ):
            _call('autonomous_campus_setup')
        elif any(w in p_lower for w in ('create subject', 'new subject', 'add subject')):
            name = _extract_quoted_or_tail(prompt, 'Data Structures & Algorithms')
            _call('autonomous_create_academic_entity', entity_type='subject', name=name)
        elif any(w in p_lower for w in ('create department', 'new department', 'add department')):
            name = _extract_quoted_or_tail(prompt, 'Computer Science & Engineering')
            _call('autonomous_create_academic_entity', entity_type='department', name=name)
        elif any(w in p_lower for w in ('create batch', 'new batch', 'add batch')):
            name = _extract_quoted_or_tail(prompt, 'CSE 2nd Year Batch')
            _call('autonomous_create_academic_entity', entity_type='batch', name=name)
        elif any(w in p_lower for w in ('create section', 'new section', 'add section')):
            name = _extract_quoted_or_tail(prompt, 'A')
            _call('autonomous_create_academic_entity', entity_type='section', name=name)
        elif any(w in p_lower for w in ('enroll', 'assign subject', 'assign instructor')):
            _call('autonomous_assign_and_enroll', enroll_students=('enroll' in p_lower))
        elif any(w in p_lower for w in ('/schedule', 'schedule session', 'recurring', 'plan session', 'timetable')):
            _call('autonomous_schedule_sessions', recurring_weeks=6 if 'recurring' in p_lower or 'weekly' in p_lower else 0)
        elif any(w in p_lower for w in ('sort', 'sorting', 'dorting')) and any(w in p_lower for w in ('creat', 'crea', 'build', 'topic', 'page', 'experiment', 'simulat', 'datastruc')):
            _call('build_complete_experiment_suite', title='Sorting Algorithms', subject_name='Data Structures', algorithm_pages=['Bubble Sort', 'Selection Sort', 'Insertion Sort', 'Merge Sort', 'Quick Sort'])
        elif any(w in p_lower for w in ('create experiment', 'build experiment', '/experiment', 'crea experiment', 'crea teh', 'create topic')):
            title = _extract_quoted_or_tail(prompt, 'Algorithms & Data Structures Lab')
            _call('build_complete_experiment_suite', title=title)
        elif any(w in p_lower for w in ('audit', 'security')):
            _call('query_audit_and_security_logs')
        elif any(w in p_lower for w in ('catalog', 'unassigned', 'hierarchy')):
            _call('query_academic_catalog_deep')
        elif any(w in p_lower for w in ('report', 'naac', 'nba', '/report')):
            _call('generate_executive_report', report_type='naac_nba_accreditation' if 'naac' in p_lower or 'nba' in p_lower else 'executive_summary')
        else:
            _call('query_analytics_domain', domain='grades')

    elif role_key == 'student':
        if any(w in p_lower for w in ('/collection', 'playlist', 'collection', 'bookmark')):
            _call('build_personal_problem_collection')
        elif any(w in p_lower for w in ('/playground', 'playground', 'workspace')):
            _call('create_saved_playground_workspace')
        elif any(w in p_lower for w in ('/studyplan', 'study plan', 'roadmap', 'enroll plan')):
            _call('enroll_and_build_study_plan')
        elif any(w in p_lower for w in ('study sheet', 'flashcard', 'cheat sheet', 'viva prep')):
            _call('generate_study_sheet')
        elif any(w in p_lower for w in ('grade', 'score', 'mark', 'gpa')):
            _call('get_my_grades')
        elif any(w in p_lower for w in ('attend', 'absent', 'present', '75%')):
            _call('get_my_attendance')
        elif any(w in p_lower for w in ('spark', 'rank', 'skill', 'weak', 'mastery')):
            _call('get_my_spark_and_skills')
        elif any(w in p_lower for w in ('upcoming', 'exam', 'schedule', 'deadline')):
            _call('get_upcoming_sessions_and_exams')
        elif any(w in p_lower for w in ('hint', 'stuck', 'clue')):
            _call('get_socratic_hint')
        else:
            _call('analyze_my_code')

    elif role_key == 'platform_user':
        if any(w in p_lower for w in ('/collection', 'playlist', 'collection', 'bookmark')):
            _call('build_personal_problem_collection')
        elif any(w in p_lower for w in ('/playground', 'playground', 'workspace')):
            _call('create_saved_playground_workspace')
        elif any(w in p_lower for w in ('/studyplan', 'study plan', 'roadmap')):
            _call('enroll_and_build_study_plan')
        elif any(w in p_lower for w in ('recommend', 'next', 'problem')):
            _call('recommend_next_problem')
        else:
            _call('get_my_practice_stats')

    elif role_key == 'admin':
        if any(w in p_lower for w in ('/onboard', 'onboard', 'provision', 'create institution')):
            name = _extract_quoted_or_tail(prompt, 'NIT Autonomous Campus')
            _call('autonomous_onboard_institution', name=name)
        elif any(w in p_lower for w in ('platform course', 'create course', 'contest')):
            title = _extract_quoted_or_tail(prompt, 'Mastering Python & Algorithms')
            _call('autonomous_create_platform_course_or_contest', entity_type='contest' if 'contest' in p_lower else 'platform_course', title=title)
        elif any(w in p_lower for w in ('list institution', 'institutions', 'tenants')):
            _call('list_platform_institutions')
        else:
            _call('get_platform_overview')

    return executed


# ══════════════════════════════════════════════════════════════════════════════
# 7. Unified Claw-Code Multi-Turn Execution Loop (`run_harness_turn`)
# ══════════════════════════════════════════════════════════════════════════════

def run_harness_turn(prompt: str, institution_id: str, user_id: str, role: str,
                     mode: str = 'chat', doc_id: str = None, web_search_enabled: bool = False,
                     chat_history: list = None, page_context: dict = None,
                     compact_memory: str = None, session_id: str = None) -> dict:
    """
    Executes a complete Claw-Code multi-turn autonomous loop (`max_turns=6`) for any of the 5 roles:
    - Intercepts `/slash` commands via `match_slash_command()`
    - Chains up to 6 sequential read/create/setup/authoring tools in a single turn
    - Auto-executes all non-destructive tools immediately in the database
    - Emits real-time Socket.IO events (`ai:tool_start`, `ai:tool_result`, `ai:token`, `ai:done`)
    - Enforces deterministic Socratic output filtering for `student` and `platform_user`
    """
    # Live role context (learner memory / teaching / platform snapshot)
    from app.ai.learner_context import build_role_context
    try:
        role_context_block = build_role_context(current_user)
    except Exception:
        role_context_block = ''
    role_key = (role or 'student').lower().strip()
    g.user_id = user_id
    g.institution_id = institution_id
    g.role = role_key
    g.harness_page_context = page_context or {}

    user_obj = User.query.get(user_id) if user_id else None
    inst_obj = Institution.query.get(institution_id) if institution_id else None

    role_tools = get_tools_for_role(role_key)

    # Check if prompt starts with a slash command (/setup, /experiment, /exam, /schedule, /collection, etc.)
    matched_cmd, effective_prompt = match_slash_command(prompt, role_key)

    catalog_str = select_routed_tool_catalog(
        role_tools=role_tools,
        role_key=role_key,
        mode=mode,
        prompt=effective_prompt,
        page_context=page_context,
        max_catalog_tools=18,
    )

    tool_calls = []
    usage_summary = UsageSummary()
    pending_action = None
    generated_doc = None
    sources = []
    pre_observations = []

    # 1. Optional Web Search
    if web_search_enabled:
        _emit_ai_event(user_id, 'ai:tool_start', {'tool': 'search_web', 'session_id': session_id})
        web_res = run_web_search(effective_prompt, max_results=4)
        tool_calls.append({'tool': 'search_web', 'args': {'query': effective_prompt}, 'result': web_res})
        sources.extend([
            {'type': 'web', 'title': w.get('title'), 'url': w.get('url'), 'snippet': w.get('snippet')}
            for w in web_res if isinstance(w, dict) and w.get('url')
        ])
        pre_observations.append(f"Web Search Results:\n{_truncate(web_res)}")
        _emit_ai_event(user_id, 'ai:tool_result', {'tool': 'search_web', 'session_id': session_id})

    # 2. Scoped Multi-Format RAG Search
    rag_chunks = []
    p_lower = effective_prompt.lower()
    if doc_id or any(k in p_lower for k in ('pdf', 'document', 'syllabus', 'manual', 'handout', 'uploaded', 'policy', 'notes')):
        rag_chunks = search_scoped_documents(institution_id, user_id, role_key, effective_prompt, doc_id=doc_id, top_k=4)
        if rag_chunks:
            tool_calls.append({'tool': 'search_uploaded_documents', 'args': {'query': effective_prompt, 'doc_id': doc_id}, 'result': rag_chunks})
            sources.extend([
                {'type': 'document', 'filename': c.get('filename'), 'page': c.get('page'), 'text': c.get('text')}
                for c in rag_chunks
            ])
            pre_observations.append(f"Document RAG Context:\n{_truncate(rag_chunks)}")

    # 3. Build 6-Layer Planner Prompt
    planner_sys = build_six_layer_prompt(
        role=role_key,
        mode=mode,
        user=user_obj,
        inst=inst_obj,
        page_context=page_context,
        compact_memory=compact_memory,
        rag_chunks=rag_chunks,
        catalog_str=catalog_str,
        is_planner=True,
        role_context=role_context_block,
    )

    planner_history = [{'role': 'system', 'content': planner_sys}]
    if chat_history:
        for m in chat_history[-5:]:
            r = 'user' if m.get('role') == 'user' else 'assistant'
            planner_history.append({'role': r, 'content': (m.get('content') or m.get('text') or '')[:1200]})

    user_content = f"User Request ({role_key}): {effective_prompt}"
    if matched_cmd:
        user_content += f"\n(Matched Slash Command `{matched_cmd.name}` -> Invoke `{matched_cmd.target_tool}`)"
    if pre_observations:
        user_content += "\n\nPre-retrieved observations:\n" + "\n\n".join(pre_observations)
    planner_history.append({'role': 'user', 'content': user_content})

    gate_feature = 'chat' if role_key in ('student', 'platform_user') else 'agent'

    # 4. Claw-Code Multi-Turn QueryEngine Loop (Up to 6 Chained Tool Calls!)
    for _turn in range(ENGINE_CONFIG.max_turns):
        if usage_summary.total_tokens >= ENGINE_CONFIG.max_budget_tokens:
            break

        text, _model, is_fallback, usage = guarded_ai_call(
            gate_feature,
            institution_id=institution_id,
            user_id=user_id,
            role=role_key,
            messages=planner_history,
            pool='quality',
            max_tokens=2048,
            temperature=0.1,
        )
        usage_summary = usage_summary.add(usage)

        if is_fallback or not text:
            break

        plan = _extract_json(text)
        tool_name = plan.get('tool')
        if not tool_name:
            break

        entry = role_tools.get(tool_name)
        if not entry:
            planner_history.append({'role': 'assistant', 'content': text})
            planner_history.append({'role': 'user', 'content': f"Unknown tool '{tool_name}'. Choose from the routed tools list."})
            continue

        fn = entry[0]
        args = plan.get('args') if isinstance(plan.get('args'), dict) else {}

        _emit_ai_event(user_id, 'ai:tool_start', {'tool': tool_name, 'args': args, 'session_id': session_id})
        try:
            result = fn(**args)
            call_record = {'tool': tool_name, 'args': args, 'result': result}

            if tool_name in CONFIRM_ACTION_NAMES or (isinstance(result, dict) and result.get('requires_confirmation')):
                pending_action = result.get('draft') or result
                _emit_ai_event(user_id, 'ai:confirm_action', {'pending_action': pending_action, 'session_id': session_id})
                _emit_ai_event(user_id, 'ai:action_card', {'pending_action': pending_action, 'session_id': session_id})

            if isinstance(result, dict) and result.get('generated_document'):
                generated_doc = result.get('generated_document')

            if tool_name == 'search_uploaded_documents' and isinstance(result, dict):
                for src in result.get('sources', []):
                    sources.append({'type': 'document', 'filename': src.get('filename'), 'page': src.get('page'), 'text': src.get('text')})
            elif tool_name == 'search_web' and isinstance(result, dict):
                for w in result.get('results', []):
                    sources.append({'type': 'web', 'title': w.get('title'), 'url': w.get('url'), 'snippet': w.get('snippet')})
        except AgentToolError as exc:
            call_record = {'tool': tool_name, 'args': args, 'error': str(exc)}
        except Exception as exc:
            log.warning("[harness.tool] %s failed: %s", tool_name, exc)
            call_record = {'tool': tool_name, 'args': args, 'error': f"Tool error: {str(exc)}"}

        _emit_ai_event(user_id, 'ai:tool_result', {'tool': tool_name, 'session_id': session_id})
        tool_calls.append(call_record)

        planner_history.append({'role': 'assistant', 'content': text})
        planner_history.append({
            'role': 'user',
            'content': (
                f"Observation from {tool_name}:\n{_truncate(call_record.get('result') or call_record.get('error'))}\n\n"
                "If additional steps are needed to fulfill the user's request, call the next tool; "
                "otherwise reply {\"tool\": null, \"reason\": \"done\"}."
            ),
        })

        # Only break early if a destructive operation requires user confirmation
        if pending_action:
            break

    # 5. Deterministic Autonomous Fallback if LLM was rate-limited or skipped tool execution
    if not tool_calls:
        fallback_calls = _run_deterministic_autonomous_fallback(effective_prompt, role_key, role_tools)
        for fc in fallback_calls:
            tool_calls.append(fc)
            res_obj = fc.get('result')
            if isinstance(res_obj, dict) and res_obj.get('generated_document'):
                generated_doc = res_obj.get('generated_document')
            if isinstance(res_obj, dict) and res_obj.get('requires_confirmation'):
                pending_action = res_obj.get('draft') or res_obj

    # 6. Synthesis Step
    obs_summary = '\n\n'.join(
        f"Tool '{c['tool']}' → " + _truncate(c.get('result') if 'result' in c else {'error': c.get('error')})
        for c in tool_calls
    )
    if pre_observations and not obs_summary:
        obs_summary = "\n\n".join(pre_observations)

    answer_sys = build_six_layer_prompt(
        role=role_key,
        mode=mode,
        user=user_obj,
        inst=inst_obj,
        page_context=page_context,
        compact_memory=compact_memory,
        rag_chunks=rag_chunks,
        is_planner=False,
        role_context=role_context_block,
    )

    answer, _model, is_fallback, usage = guarded_ai_call(
        gate_feature,
        institution_id=institution_id,
        user_id=user_id,
        role=role_key,
        messages=[
            {'role': 'system', 'content': answer_sys},
            {
                'role': 'user',
                'content': (
                    f"User Prompt: {prompt}\n\n"
                    f"Executed Tools & Database Results:\n{obs_summary or '(No database tools invoked)'}"
                ),
            },
        ],
        pool='quality',
        max_tokens=1800,
        temperature=0.25,
    )
    usage_summary = usage_summary.add(usage)

    if is_fallback or not answer:
        # Build a rich, deterministic summary from the executed tool results!
        msgs = []
        for c in tool_calls:
            r = c.get('result')
            if isinstance(r, dict) and r.get('message'):
                msgs.append(f"- **{c['tool']}**: {r['message']}")
        if msgs:
            answer = "### Autonomous Execution Summary\n\n" + "\n".join(msgs)
        elif pending_action:
            answer = "I've prepared the confirmation card for this destructive operation below."
        elif generated_doc:
            answer = f"I've generated **{generated_doc.get('title')}** for you. Click the download button below to save it."
        else:
            answer = "Here are the latest results from your NexLab workspace."

    # 7. Enforce Deterministic Socratic Output Guardrail for Student & Platform User
    answer, socratic_filtered = enforce_socratic_guardrail(answer, role_key)

    # Sanitize hallucinated "Confirm & Apply" promises if no actual pending_action card exists
    if not pending_action and answer:
        answer = re.sub(r'(?i)(?:please\s+)?confirm\s+you(?:\'d| would)\s+like\s+me\s+to\s+run\s+this\s+query[^\n]*\n?', '', answer)
        answer = re.sub(r'(?i)[*_\s]*once\s+you\s+click\s+.*confirm\s*(?:&|and)\s*apply[^\n]*\n?', '', answer)
        answer = re.sub(r'(?i)[*_\s]*(?:please\s+)?click\s+.*confirm\s*(?:&|and)\s*apply[^\n]*\n?', '', answer)
        answer = re.sub(r'(?i)[*_\s]*on\s+the\s+action\s+card\s+below[^\n]*\n?', '', answer)
        answer = answer.strip()

    _emit_ai_event(user_id, 'ai:token', {'session_id': session_id, 'delta': answer})
    _emit_ai_event(user_id, 'ai:done', {
        'session_id': session_id,
        'total_tokens': usage_summary.total_tokens,
        'socratic_filtered': socratic_filtered,
    })

    return {
        'status': 'pending_confirm' if pending_action else 'completed',
        'answer': answer,
        'tool_calls': tool_calls,
        'total_tokens': usage_summary.total_tokens,
        'pending_action': pending_action,
        'generated_doc': generated_doc,
        'socratic_filtered': socratic_filtered,
        'sources': sources,
    }
